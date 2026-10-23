"""
ReelLore AI — Automated Academic Report Generator
====================================================
Generates a complete OJT/Coursera Project Report (.docx)
following strict university guidelines using python-docx.
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


# ── Helper Functions ────────────────────────────────────────────────

def set_cell_shading(cell, color_hex: str):
    """Apply background shading to a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def add_heading_styled(doc, text, level=1):
    """Add a heading with consistent styling."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    return heading


def add_para(doc, text, bold=False, italic=False, size=11, alignment=None, space_after=6):
    """Add a paragraph with formatting."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if alignment:
        p.alignment = alignment
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_code_block(doc, code_text, title=""):
    """Add a formatted code block."""
    if title:
        add_para(doc, title, bold=True, size=10)
    p = doc.add_paragraph()
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x20, 0x20, 0x30)
    # Light gray background via paragraph shading
    pPr = p._p.get_or_add_pPr()
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F0F5" w:val="clear"/>')
    pPr.append(shading)
    return p


def add_screenshot_placeholder(doc, caption):
    """Add a bordered placeholder box for screenshots."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"\n\n{caption}\n\n")
    run.bold = True
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x99)
    set_cell_shading(cell, "F5F5FA")
    # Border
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="8" w:color="8032FC"/>'
        f'<w:left w:val="single" w:sz="8" w:color="8032FC"/>'
        f'<w:bottom w:val="single" w:sz="8" w:color="8032FC"/>'
        f'<w:right w:val="single" w:sz="8" w:color="8032FC"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)
    doc.add_paragraph()


def read_source_file(filepath):
    """Read source code from a project file."""
    base_dir = os.path.dirname(os.path.abspath(__file__))
    full_path = os.path.join(base_dir, filepath)
    try:
        with open(full_path, "r", encoding="utf-8") as f:
            return f.read()
    except FileNotFoundError:
        return f"# File not found: {filepath}"


# ── Report Generation ──────────────────────────────────────────────

def generate_report():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(2.54)

    # Default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    # ════════════════════════════════════════════════════════════════
    # COVER PAGE
    # ════════════════════════════════════════════════════════════════
    for _ in range(4):
        doc.add_paragraph()

    add_para(doc, "SVKM's Dwarkadas J. Sanghvi College of Engineering",
             bold=True, size=14, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Department of Computer Engineering",
             bold=True, size=12, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "(Honors in Data Science)",
             italic=True, size=11, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    add_para(doc, "OJT / Coursera Project Report",
             bold=True, size=13, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    # Project title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Design and Implementation of ReelLore AI:\n"
        "An Autonomous Multi-Agent Game Content\n"
        "Curation System using CrewAI"
    )
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x80, 0x32, 0xFC)
    p.paragraph_format.space_after = Pt(36)

    # Student details
    details = [
        ("Student Name", "Shrav Jitendra Siddhpura"),
        ("SAP ID", "60004230236"),
        ("Roll Number", "C143"),
        ("Programme", "B.Tech in Computer Engineering (Honors in Data Science)"),
        ("Institute", "DJSCE, Mumbai"),
        ("Academic Year", "2026–2027"),
    ]
    table = doc.add_table(rows=len(details), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, value) in enumerate(details):
        cell_l = table.cell(i, 0)
        cell_r = table.cell(i, 1)
        cell_l.text = label
        cell_r.text = value
        for cell in [cell_l, cell_r]:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(11)
                    run.font.name = "Calibri"
            set_cell_shading(cell, "F5F5FA")
        cell_l.paragraphs[0].runs[0].bold = True

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # TABLE OF CONTENTS
    # ════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "Table of Contents", level=1)

    toc_entries = [
        ("1", "Introduction to Agentic AI Systems"),
        ("1.1", "Autonomous Agents and the ReAct Paradigm"),
        ("1.2", "Multi-Agent Systems (MAS)"),
        ("1.3", "CrewAI Framework"),
        ("2", "Problem Definition"),
        ("2.1", "Challenges in Short-Form Gaming Content"),
        ("2.2", "Problem Statement"),
        ("2.3", "Objectives"),
        ("3", "System Design & Multi-Agent Architecture"),
        ("3.1", "Agent Specifications"),
        ("3.2", "Custom Tool Specifications"),
        ("3.3", "LLM Configuration"),
        ("3.4", "Inter-Agent Context Handoff Flow"),
        ("4", "Development & Implementation Details"),
        ("4.1", "Custom Tools Implementation"),
        ("4.2", "Agent Definitions"),
        ("4.3", "Task Definitions"),
        ("4.4", "Crew Orchestration"),
        ("4.5", "Screenshots"),
        ("5", "Deployment Architecture & Proof of Multi-Agent Execution"),
        ("5.1", "Streamlit Community Cloud Deployment"),
        ("5.2", "Secret Management & API Integration"),
        ("5.3", "Callback Telemetry as Proof"),
        ("6", "Conclusion & Implementation Proof Links"),
        ("6.1", "Summary of Achievements"),
        ("6.2", "Verification Table"),
    ]

    toc_table = doc.add_table(rows=len(toc_entries), cols=2)
    toc_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (num, title) in enumerate(toc_entries):
        toc_table.cell(i, 0).text = num
        toc_table.cell(i, 1).text = title
        for cell in [toc_table.cell(i, 0), toc_table.cell(i, 1)]:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.name = "Calibri"
        # Bold for chapter-level entries
        if "." not in num:
            for cell in [toc_table.cell(i, 0), toc_table.cell(i, 1)]:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # CHAPTER 1: Introduction to Agentic AI Systems
    # ════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "Chapter 1: Introduction to Agentic AI Systems", level=1)

    add_heading_styled(doc, "1.1 Autonomous Agents and the ReAct Paradigm", level=2)
    add_para(doc, (
        "The field of Artificial Intelligence has witnessed a paradigm shift from static, single-turn "
        "language model interactions toward autonomous agent systems capable of sustained, goal-directed "
        "behavior. An AI agent, in the context of Large Language Models (LLMs), is a software entity "
        "that receives a high-level objective, decomposes it into sub-tasks, selects and invokes "
        "external tools, interprets intermediate results, and iteratively refines its approach until "
        "the goal is accomplished — all without continuous human intervention."
    ))
    add_para(doc, (
        "The ReAct (Reasoning + Acting) paradigm, introduced by Yao et al. (2023), formalizes this "
        "behavior by interleaving chain-of-thought reasoning traces with concrete tool-use actions. "
        "At each step the agent: (1) Observes the current state and any tool outputs, (2) Thinks — "
        "producing an internal reasoning trace that analyzes progress and plans the next move, and "
        "(3) Acts — choosing a tool to invoke or producing a final answer. This Thought → Action → "
        "Observation loop enables agents to ground their reasoning in real-world data, self-correct "
        "errors, and handle tasks that a monolithic LLM call cannot reliably solve."
    ))

    add_heading_styled(doc, "1.2 Multi-Agent Systems (MAS)", level=2)
    add_para(doc, (
        "While a single agent can accomplish focused tasks, complex real-world workflows benefit from "
        "Multi-Agent Systems where multiple specialized agents collaborate. Each agent is assigned a "
        "distinct role, skill set (tools), and personality (backstory), enabling division of labor and "
        "expert-level performance in each sub-domain. MAS architectures offer several advantages over "
        "monolithic agents: (a) Specialization — each agent masters a narrow domain; (b) Modularity — "
        "agents can be added, removed, or swapped without rewriting the pipeline; (c) Scalability — "
        "parallel or sequential execution across agents; and (d) Interpretability — the output of each "
        "agent can be independently inspected for quality and correctness."
    ))

    add_heading_styled(doc, "1.3 CrewAI Framework", level=2)
    add_para(doc, (
        "CrewAI is a Python framework that operationalizes the multi-agent paradigm through a "
        "role-playing metaphor. Developers define Agents (with roles, goals, backstories, and tools), "
        "Tasks (with descriptions, expected outputs, and context dependencies), and a Crew (that "
        "orchestrates execution in sequential or hierarchical order). CrewAI integrates with multiple "
        "LLM providers via LiteLLM, supports tool callbacks for observability, and provides structured "
        "output parsing. In this project, CrewAI is used to orchestrate a 4-agent sequential pipeline "
        "for automated game content research, scriptwriting, and production."
    ))

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # CHAPTER 2: Problem Definition
    # ════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "Chapter 2: Problem Definition", level=1)

    add_heading_styled(doc, "2.1 Challenges in Short-Form Gaming Content", level=2)
    add_para(doc, (
        "The gaming content creation industry, particularly for platforms like Instagram Reels, TikTok, "
        "and YouTube Shorts, faces several challenges that make high-quality, fact-based content "
        "production inefficient and time-consuming:"
    ))

    challenges = [
        "Fragmented Information Sources: Game lore, Easter eggs, and hidden details are scattered "
        "across multiple wikis (Fandom, IGN, GameFAQs), developer interviews, and data-mining "
        "communities with no unified aggregation point.",
        "Unstructured Community Data: Reddit threads and community forums contain valuable bug "
        "reports, dev stories, and Easter egg discoveries, but in deeply nested, unstructured "
        "comment threads that require significant manual curation.",
        "Raw Patch Notes: Steam and developer update logs contain fascinating engineering stories "
        "about bug fixes and workarounds, but are written in technical jargon inaccessible to "
        "casual gaming audiences.",
        "Script-to-Production Gap: Even when compelling stories are found, transforming raw research "
        "into a precisely timed, scroll-stopping Reel script with B-roll cues, voiceover tags, and "
        "SEO captions requires a separate, specialized skill set.",
        "Manual Bottleneck: Content creators currently perform all four roles — researcher, writer, "
        "scriptwriter, and producer — manually, limiting output to 2-3 Reels per week.",
    ]
    for c in challenges:
        add_para(doc, f"• {c}", size=10)

    add_heading_styled(doc, "2.2 Problem Statement", level=2)
    add_para(doc, (
        "To design and implement an autonomous multi-agent system that accepts a video game title "
        "as input and produces a complete, production-ready Instagram Reel content package — including "
        "researched lore findings, developer anecdotes, three 45-60 second Reel scripts, and a full "
        "production sheet with B-roll cues, voiceover tone tags, and SEO-optimized captions — by "
        "orchestrating four specialized AI agents through CrewAI's sequential pipeline."
    ), bold=True)

    add_heading_styled(doc, "2.3 Objectives", level=2)
    objectives = [
        "Automate game lore and trivia research using web search APIs.",
        "Mine community platforms (Reddit) and developer channels (Steam) for bug stories and patch notes.",
        "Generate multiple production-ready Instagram Reel scripts with precise timing structures.",
        "Produce comprehensive production sheets including visual cues, voiceover tags, and Instagram SEO.",
        "Provide transparent multi-agent execution telemetry for verifiability.",
        "Deploy as a self-service web application accessible via Streamlit Community Cloud.",
    ]
    for i, obj in enumerate(objectives, 1):
        add_para(doc, f"{i}. {obj}", size=10)

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # CHAPTER 3: System Design & Multi-Agent Architecture
    # ════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "Chapter 3: System Design & Multi-Agent Architecture", level=1)

    add_heading_styled(doc, "3.1 Agent Specifications", level=2)
    add_para(doc, "The system comprises four sequential agents, each specialized for a stage of the content pipeline:")

    # Agent table
    agent_table = doc.add_table(rows=5, cols=4)
    agent_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    agent_table.style = "Table Grid"
    headers = ["Agent", "Role", "Tools", "Output"]
    for i, h in enumerate(headers):
        cell = agent_table.cell(0, i)
        cell.text = h
        set_cell_shading(cell, "8032FC")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.bold = True
                run.font.size = Pt(9)

    agent_data = [
        ["Agent 1", "Lore_And_Trivia_Scout", "SerperGameSearchTool", "01_lore_findings.md"],
        ["Agent 2", "DevLog_And_Bug_Archaeologist", "RedditMinerTool, SteamPatchNotesTool", "02_dev_and_bugs.md"],
        ["Agent 3", "Viral_Reel_Scriptwriter", "None (uses context)", "03_reel_scripts.md"],
        ["Agent 4", "Production_And_SEO_Director", "None (uses context)", "04_production_package.md"],
    ]
    for row_idx, row_data in enumerate(agent_data, 1):
        for col_idx, val in enumerate(row_data):
            cell = agent_table.cell(row_idx, col_idx)
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()

    add_heading_styled(doc, "3.2 Custom Tool Specifications", level=2)

    tool_table = doc.add_table(rows=4, cols=3)
    tool_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tool_table.style = "Table Grid"
    tool_headers = ["Tool Name", "API Endpoint", "Purpose"]
    for i, h in enumerate(tool_headers):
        cell = tool_table.cell(0, i)
        cell.text = h
        set_cell_shading(cell, "8032FC")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.bold = True
                run.font.size = Pt(9)

    tool_data = [
        ["SerperGameSearchTool", "google.serper.dev/search", "Google web search for game lore and developer interviews"],
        ["RedditMinerTool", "reddit.com/search.json", "Community mining for bugs, Easter eggs, dev stories (with Serper fallback)"],
        ["SteamPatchNotesTool", "store.steampowered.com + api.steampowered.com", "Fetch real developer patch notes and update logs"],
    ]
    for row_idx, row_data in enumerate(tool_data, 1):
        for col_idx, val in enumerate(row_data):
            cell = tool_table.cell(row_idx, col_idx)
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()

    add_heading_styled(doc, "3.3 LLM Configuration", level=2)
    add_para(doc, (
        "All agents utilize the Groq LPU Inference Engine for high-speed LLM inference. "
        "The primary model is LLaMA 3.3 70B Versatile (groq/llama-3.3-70b-versatile), with "
        "LLaMA 3.1 8B Instant (groq/llama-3.1-8b-instant) available as a fast fallback option. "
        "Configuration: temperature=0.7, max_iter=3 (per agent), max_rpm=15 (rate-limit safety). "
        "API authentication is via GROQ_API_KEY loaded from environment variables."
    ))

    add_heading_styled(doc, "3.4 Inter-Agent Context Handoff Flow", level=2)
    add_para(doc, (
        "The crew follows a strictly sequential process:\n\n"
        "1. Agent 1 (Lore Scout) executes independently → produces lore findings\n"
        "2. Agent 2 (Bug Archaeologist) executes independently → produces dev-log stories\n"
        "3. Agent 3 (Scriptwriter) receives context from Tasks 1 & 2 → writes 3 Reel scripts\n"
        "4. Agent 4 (Production Director) receives context from Task 3 → creates production package\n\n"
        "This ensures each downstream agent builds upon verified upstream research, maintaining "
        "factual accuracy while adding creative value."
    ))

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # CHAPTER 4: Development & Implementation Details
    # ════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "Chapter 4: Development & Implementation Details", level=1)

    add_heading_styled(doc, "4.1 Custom Tools Implementation (src/tools.py)", level=2)
    tools_code = read_source_file("src/tools.py")
    add_code_block(doc, tools_code, "Source: src/tools.py")
    doc.add_paragraph()

    add_heading_styled(doc, "4.2 Agent Definitions (src/agents.py)", level=2)
    agents_code = read_source_file("src/agents.py")
    add_code_block(doc, agents_code, "Source: src/agents.py")
    doc.add_paragraph()

    add_heading_styled(doc, "4.3 Task Definitions (src/tasks.py)", level=2)
    tasks_code = read_source_file("src/tasks.py")
    add_code_block(doc, tasks_code, "Source: src/tasks.py")
    doc.add_paragraph()

    add_heading_styled(doc, "4.4 Crew Orchestration (src/crew.py)", level=2)
    crew_code = read_source_file("src/crew.py")
    add_code_block(doc, crew_code, "Source: src/crew.py")

    doc.add_page_break()

    add_heading_styled(doc, "4.5 Application Screenshots", level=2)
    add_para(doc, "The following screenshots demonstrate the working application and multi-agent execution:")

    screenshots = [
        "[INSERT SCREENSHOT 1: Streamlit Main UI & Game Input]",
        "[INSERT SCREENSHOT 2: Proof of 4 Individual Agent Outputs Tab]",
        "[INSERT SCREENSHOT 3: Live Agent Execution Trace & Tool Calls]",
        "[INSERT SCREENSHOT 4: Multi-Agent Architecture Diagram Tab]",
    ]
    for caption in screenshots:
        add_screenshot_placeholder(doc, caption)

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # CHAPTER 5: Deployment Architecture & Proof of Multi-Agent Execution
    # ════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "Chapter 5: Deployment Architecture & Proof of Multi-Agent Execution", level=1)

    add_heading_styled(doc, "5.1 Streamlit Community Cloud Deployment", level=2)
    add_para(doc, (
        "ReelLore AI is deployed as a web application on Streamlit Community Cloud, which provides "
        "a free hosting environment for Python-based data applications. The deployment workflow involves: "
        "(1) Pushing the project repository to GitHub, (2) Connecting the repository to Streamlit Cloud, "
        "(3) Configuring API secrets in the Streamlit Cloud dashboard, and (4) Setting the entry point "
        "as app.py. The .streamlit/config.toml file ensures the dark theme and branding are preserved "
        "across environments."
    ))

    add_heading_styled(doc, "5.2 Secret Management & API Integration", level=2)
    add_para(doc, (
        "The application integrates with three external APIs, managed through a dual secret-loading strategy:"
    ))

    secret_table = doc.add_table(rows=4, cols=3)
    secret_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    secret_table.style = "Table Grid"
    sec_headers = ["Secret / API", "Local Development", "Cloud Deployment"]
    for i, h in enumerate(sec_headers):
        cell = secret_table.cell(0, i)
        cell.text = h
        set_cell_shading(cell, "8032FC")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.bold = True
                run.font.size = Pt(9)

    sec_data = [
        ["GROQ_API_KEY", ".env file (python-dotenv)", "st.secrets (Streamlit Cloud)"],
        ["SERPER_API_KEY", ".env file (python-dotenv)", "st.secrets (Streamlit Cloud)"],
        ["Steam Web API", "No key required (public)", "No key required (public)"],
    ]
    for row_idx, row_data in enumerate(sec_data, 1):
        for col_idx, val in enumerate(row_data):
            cell = secret_table.cell(row_idx, col_idx)
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()

    add_heading_styled(doc, "5.3 Callback Telemetry as Proof of Multi-Agent Execution", level=2)
    add_para(doc, (
        "A critical requirement for academic verification is demonstrating that multiple distinct agents "
        "executed independently. ReelLore AI implements this through CrewAI's callback system:"
    ))
    add_para(doc, (
        "• step_callback: Fires after every agent reasoning step, logging the agent's role, the tool "
        "called (if any), and a snippet of the reasoning trace. This proves that each agent independently "
        "invoked its assigned tools (e.g., Agent 1 called SerperGameSearchTool, Agent 2 called "
        "RedditMinerTool and SteamPatchNotesTool)."
    ), size=10)
    add_para(doc, (
        "• task_callback: Fires when a task completes, logging the task output summary and timestamp. "
        "This proves that all four tasks executed sequentially and produced distinct outputs."
    ), size=10)
    add_para(doc, (
        "• Individual Output Files: Each agent's output is saved to a separate file in the outputs/ "
        "directory (01_lore_findings.md through 04_production_package.md), providing file-level proof "
        "of multi-agent execution."
    ), size=10)
    add_para(doc, (
        "• UI Verification: Tab 2 of the Streamlit interface displays each agent's output in separate "
        "expanders with word counts and timestamps. Tab 3 shows the complete execution trace log."
    ), size=10)

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # CHAPTER 6: Conclusion & Implementation Proof Links
    # ════════════════════════════════════════════════════════════════
    add_heading_styled(doc, "Chapter 6: Conclusion & Implementation Proof Links", level=1)

    add_heading_styled(doc, "6.1 Summary of Achievements", level=2)
    add_para(doc, (
        "This project successfully demonstrates the design, implementation, and deployment of ReelLore AI, "
        "an autonomous multi-agent system for game content curation. Key achievements include:"
    ))
    achievements = [
        "Designed and implemented a 4-agent sequential pipeline using CrewAI for automated game content research and production.",
        "Integrated three external data sources (Google Serper, Reddit JSON API, Steam Web API) through custom tool implementations.",
        "Generated production-ready Instagram Reel scripts with precise timing structures (Hook, Body, Climax/CTA).",
        "Built a premium Streamlit web application with dark theme, live execution telemetry, and multi-agent proof tabs.",
        "Deployed the application to Streamlit Community Cloud with secure API key management.",
        "Implemented callback-based telemetry for transparent, verifiable multi-agent execution proof.",
        "Used Groq LPU Inference Engine for high-speed, cost-effective LLM inference (LLaMA 3.3 70B).",
    ]
    for a in achievements:
        add_para(doc, f"• {a}", size=10)

    add_heading_styled(doc, "6.2 Verification Table", level=2)
    add_para(doc, "The following table provides all links necessary for implementation verification:")

    verify_table = doc.add_table(rows=5, cols=2)
    verify_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    verify_table.style = "Table Grid"

    verify_data = [
        ["Verification Item", "Details"],
        ["Student Details", "Shrav Jitendra Siddhpura\nSAP ID: 60004230236 | Roll No: C143\nB.Tech Computer Engineering (Honors in Data Science)\nDJSCE, Mumbai"],
        ["Live Deployed Application URL", "[PASTE_STREAMLIT_CLOUD_URL_HERE]"],
        ["GitHub Repository URL", "[PASTE_GITHUB_REPO_URL_HERE]"],
        ["Google Drive Demo Video Link\n(Proof of Implementation)", "[PASTE_GOOGLE_DRIVE_DEMO_LINK_HERE]"],
    ]

    for row_idx, (col1, col2) in enumerate(verify_data):
        cell_l = verify_table.cell(row_idx, 0)
        cell_r = verify_table.cell(row_idx, 1)
        cell_l.text = col1
        cell_r.text = col2
        if row_idx == 0:
            for cell in [cell_l, cell_r]:
                set_cell_shading(cell, "8032FC")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.bold = True
                        run.font.size = Pt(10)
        else:
            for cell in [cell_l, cell_r]:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(10)
            cell_l.paragraphs[0].runs[0].bold = True
            # Add purple left border to verification items
            tc_pr = cell_l._tc.get_or_add_tcPr()
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="12" w:color="8032FC"/>'
                f'</w:tcBorders>'
            )
            tc_pr.append(borders)

    doc.add_paragraph()
    doc.add_paragraph()

    # Final signature line
    add_para(doc, "—— End of Report ——",
             bold=True, size=12, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # ── Save ────────────────────────────────────────────────────────
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "OJT_Coursera_Project_Report.docx")
    doc.save(output_path)
    print(f"✅ Report generated successfully: {output_path}")
    return output_path


if __name__ == "__main__":
    generate_report()
